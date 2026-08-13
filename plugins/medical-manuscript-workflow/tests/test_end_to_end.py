import copy
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from scripts.assemble_submission_package import assemble_package
from scripts.build_reviewer_responses import build_responses
from scripts.self_check import structure_issues
from scripts.validate_reviewer_coverage import coverage_issues
from scripts.validate_revision_release import release_issues
from scripts.validate_submission_package import validate_submission_package
from scripts.workflow_common import sha256_file


HASH_A = "a" * 64
HASH_B = "b" * 64


def approved_comment(comment_id: str, order: int, source: str, start: int, end: int, *, deletion: bool = False) -> dict:
    before = "Sensitivity analyses were exploratory." if deletion else "The endpoint was assessed at Week 12."
    after = "" if deletion else "The prespecified primary endpoint was assessed at Week 12."
    response = (
        "We removed the redundant sentence about the exploratory status and retained the prespecified analysis description."
        if deletion
        else "We clarified the prespecified primary endpoint in the Methods section."
    )
    return {
        "comment_id": comment_id,
        "order": order,
        "source_span": {"start": start, "end": end},
        "verbatim_comment": source[start:end],
        "core_request": "Address the Reviewer request.",
        "proposed_action": "Revise the manuscript with source-supported wording.",
        "affected_artifacts": ["Manuscript_Clean.docx", "Manuscript_Highlighted.docx"],
        "decision": {
            "state": "approved",
            "decided_at": "2026-08-06T09:00:00+08:00",
            "conditions": [],
            "conditions_met": True,
            "rationale": "The requested clarification is supported by the study records.",
        },
        "implementation": {
            "status": "verified",
            "evidence": [
                {
                    "artifact": "Manuscript_Clean.docx",
                    "before": before,
                    "after": after,
                    "anchor": "Methods > Outcomes" if not deletion else "Results > Sensitivity Analysis",
                    "sha256": HASH_B,
                }
            ],
        },
        "response": response,
        "location": {
            "status": "verified",
            "entries": [
                {
                    "artifact": "Manuscript_Clean.docx",
                    "section": "Methods > Outcomes" if not deletion else "Results > Sensitivity Analysis",
                    "page": 5,
                    "paragraph": 2,
                    "anchor": "prespecified primary endpoint" if not deletion else "Adjacent sensitivity-analysis sentence",
                }
            ],
        },
        "status": "closed",
    }


class EndToEndReviewerRevisionTests(unittest.TestCase):
    def make_ledger(self) -> dict:
        source_one = (
            "The topic is clinically relevant. Please define the primary endpoint more clearly and state whether it was prespecified. "
            "The sensitivity analysis is useful; please report its sample size and explain the missing-data assumption."
        )
        first_start = source_one.index("Please define")
        first_end = source_one.index(" The sensitivity")
        second_start = source_one.index("please report")
        second_end = len(source_one)
        reviewer_one = {
            "reviewer_id": "R1",
            "sequence": 1,
            "source_name": "Reviewer 1",
            "source_text": source_one,
            "source_sha256": HASH_A,
            "comments": [
                approved_comment("R1-C1", 1, source_one, first_start, first_end),
                approved_comment("R1-C2", 2, source_one, second_start, second_end, deletion=True),
            ],
            "non_action_segments": [
                {
                    "segment_id": "R1-N1",
                    "source_span": {"start": 0, "end": first_start},
                    "verbatim_text": source_one[0:first_start],
                    "reason": "Positive introductory assessment",
                },
                {
                    "segment_id": "R1-N2",
                    "source_span": {"start": first_end, "end": second_start},
                    "verbatim_text": source_one[first_end:second_start],
                    "reason": "Context for the second request",
                },
            ],
        }
        source_two = "Please state the analysis population in the Methods section."
        reviewer_two = {
            "reviewer_id": "R2",
            "sequence": 2,
            "source_name": "Reviewer 2",
            "source_text": source_two,
            "source_sha256": HASH_B,
            "comments": [approved_comment("R2-C1", 1, source_two, 0, len(source_two))],
            "non_action_segments": [],
        }
        return {
            "schema_version": "1.0",
            "project_id": "MMW-E2E-001",
            "manuscript_sha256": HASH_A,
            "reviewers": [reviewer_one, reviewer_two],
        }

    def make_docx(self, path: Path, text: str) -> None:
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    def test_full_reviewer_revision_release(self) -> None:
        self.assertEqual([], structure_issues())
        ledger = self.make_ledger()
        self.assertEqual([], coverage_issues(ledger))
        self.assertEqual([], release_issues(ledger))
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            formal = root / "formal"
            formal.mkdir()
            response_dir = formal / "responses"
            responses = build_responses(ledger, response_dir)

            clean = formal / "Manuscript_Clean.docx"
            self.make_docx(clean, "The prespecified primary endpoint was assessed in the defined analysis population.")
            highlighted = formal / "Manuscript_Highlighted.docx"
            highlighted.write_bytes(clean.read_bytes())
            document = Document(highlighted)
            document.paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
            document.save(highlighted)

            figure = formal / "Figure_1.png"
            figure.write_bytes(b"figure")
            table = formal / "Table_1.docx"
            self.make_docx(table, "Table 1")
            supplement = formal / "Supplementary_Table_1.docx"
            self.make_docx(supplement, "Supplementary Table 1")

            entries = [
                ("manuscript_clean", clean, "Manuscript_Clean.docx", None),
                ("manuscript_highlighted", highlighted, "Manuscript_Highlighted.docx", None),
                ("main_figure", figure, "Main_Figures_and_Tables/Figure_1.png", None),
                ("main_table", table, "Main_Figures_and_Tables/Table_1.docx", None),
                ("supplementary_table", supplement, "Supplementary_Materials/Supplementary_Table_1.docx", None),
            ]
            for index, response in enumerate(responses, start=1):
                entries.append(
                    (
                        "reviewer_response",
                        response,
                        f"Response_to_Reviewers/{response.name}",
                        f"R{index}",
                    )
                )
            policy = {
                "schema_version": "1.0",
                "workflow_type": "reviewer_revision",
                "package_name": "Submission_Package",
                "required_roles": ["manuscript_clean", "manuscript_highlighted", "reviewer_response"],
                "allowed_extensions": [".docx", ".png", ".tif", ".tiff", ".jpg", ".jpeg", ".pdf", ".xlsx"],
                "prohibited_extensions": [".json", ".md", ".csv", ".log", ".py", ".ps1", ".xml", ".tmp", ".bak"],
                "prohibited_name_patterns": ["^\\.", "~\\$", "(?i)(audit|ledger|work[-_ ]?log|execution|process|draft|backup|temp|cache)"],
                "registered_artifacts": [],
            }
            for role, source, relative, reviewer_id in entries:
                item = {
                    "role": role,
                    "source": str(source),
                    "relative_path": relative,
                    "sha256": sha256_file(source),
                }
                if reviewer_id:
                    item["reviewer_id"] = reviewer_id
                policy["registered_artifacts"].append(item)
            package = assemble_package(policy, root / "delivery")
            self.assertEqual([], validate_submission_package(package, policy))
            delivered = {path.relative_to(package).as_posix() for path in package.rglob("*") if path.is_file()}
            self.assertEqual({item["relative_path"] for item in policy["registered_artifacts"]}, delivered)
            self.assertFalse(any(Path(path).suffix.lower() in {".json", ".md", ".csv", ".log"} for path in delivered))


if __name__ == "__main__":
    unittest.main()
