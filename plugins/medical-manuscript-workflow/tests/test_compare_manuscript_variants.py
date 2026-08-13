import base64
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.compare_manuscript_variants import compare_variants


PNG_ONE = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def replace_zip_member(path: Path, member: str, transform) -> None:
    temporary = path.with_suffix(".tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == member:
                data = transform(data)
            target.writestr(item, data)
    temporary.replace(path)


class ManuscriptVariantTests(unittest.TestCase):
    def create_clean(self, directory: Path) -> Path:
        image = directory / "figure.png"
        image.write_bytes(PNG_ONE)
        document = Document()
        paragraph = document.add_paragraph("The primary endpoint improved at Week 12.")
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), "9")
        bookmark_start.set(qn("w:name"), "EndpointAnchor")
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), "9")
        paragraph._p.insert(0, bookmark_start)
        paragraph._p.append(bookmark_end)
        field_run = document.add_paragraph().add_run()
        field_begin = OxmlElement("w:fldChar")
        field_begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " REF EndpointAnchor \\h "
        field_end = OxmlElement("w:fldChar")
        field_end.set(qn("w:fldCharType"), "end")
        field_run._r.extend([field_begin, instruction, field_end])
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Group"
        table.cell(0, 1).text = "Result"
        document.add_picture(str(image))
        clean = directory / "Manuscript_Clean.docx"
        document.save(clean)
        return clean

    def create_highlighted(self, clean: Path, directory: Path) -> Path:
        highlighted = directory / "Manuscript_Highlighted.docx"
        shutil.copy2(clean, highlighted)
        document = Document(highlighted)
        document.paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
        document.save(highlighted)
        return highlighted

    def test_highlight_only_difference_passes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            clean = self.create_clean(directory)
            highlighted = self.create_highlighted(clean, directory)
            self.assertEqual([], compare_variants(clean, highlighted))

    def test_text_mismatch_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            clean = self.create_clean(directory)
            highlighted = self.create_highlighted(clean, directory)
            document = Document(highlighted)
            document.paragraphs[0].runs[0].text = "A different endpoint result."
            document.save(highlighted)
            self.assertIn("visible_text_mismatch", {issue["code"] for issue in compare_variants(clean, highlighted)})

    def test_protected_and_structural_mismatches_fail(self) -> None:
        mutation_names = ["field", "style", "drawing", "table", "custom_xml"]
        for mutation in mutation_names:
            with self.subTest(mutation=mutation), tempfile.TemporaryDirectory() as temporary:
                directory = Path(temporary)
                clean = self.create_clean(directory)
                highlighted = self.create_highlighted(clean, directory)
                if mutation == "field":
                    replace_zip_member(
                        highlighted,
                        "word/document.xml",
                        lambda data: data.replace(b" REF EndpointAnchor", b" REF EndpointAnchoZ", 1),
                    )
                elif mutation == "style":
                    document = Document(highlighted)
                    document.paragraphs[0].style = "Heading 1"
                    document.save(highlighted)
                elif mutation == "drawing":
                    with zipfile.ZipFile(highlighted, "r") as archive:
                        media_name = next(name for name in archive.namelist() if name.startswith("word/media/"))
                    replace_zip_member(highlighted, media_name, lambda data: data + b"changed")
                elif mutation == "table":
                    document = Document(highlighted)
                    document.tables[0].cell(0, 1).text = "Changed"
                    document.save(highlighted)
                elif mutation == "custom_xml":
                    replace_zip_member(highlighted, "customXml/item1.xml", lambda data: data.replace(b"APA", b"AMA"))
                codes = {issue["code"] for issue in compare_variants(clean, highlighted)}
                expected = {
                    "field": "field_mismatch",
                    "style": "style_mismatch",
                    "drawing": "drawing_mismatch",
                    "table": "table_mismatch",
                    "custom_xml": "custom_xml_mismatch",
                }[mutation]
                self.assertIn(expected, codes)

    def test_non_yellow_revision_highlight_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            clean = self.create_clean(directory)
            highlighted = self.create_highlighted(clean, directory)
            document = Document(highlighted)
            document.paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
            document.save(highlighted)
            self.assertIn("invalid_highlight_color", {issue["code"] for issue in compare_variants(clean, highlighted)})


if __name__ == "__main__":
    unittest.main()
