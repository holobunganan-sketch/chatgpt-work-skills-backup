import copy
import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from scripts.assemble_submission_package import assemble_package
from scripts.validate_submission_package import validate_submission_package
from scripts.workflow_common import sha256_file


class SubmissionPackageTests(unittest.TestCase):
    def make_docx(self, path: Path, text: str) -> None:
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    def make_response(self, path: Path, reviewer_number: int) -> None:
        document = Document()
        document.add_heading(f"Response to Reviewer {reviewer_number}", level=1)
        for heading, body in (
            ("Reviewer Comment", "Please clarify the endpoint."),
            ("Response", "We clarified the endpoint."),
            ("Location of Revision", "Methods; Page 4; endpoint definition"),
        ):
            document.add_heading(heading, level=2)
            document.add_paragraph(body)
        document.save(path)

    def make_project(self, root: Path) -> tuple[dict, Path]:
        sources = root / "formal_sources"
        sources.mkdir()
        clean = sources / "Manuscript_Clean.docx"
        highlighted = sources / "Manuscript_Highlighted.docx"
        self.make_docx(clean, "The prespecified primary endpoint improved.")
        shutil.copy2(clean, highlighted)
        highlighted_doc = Document(highlighted)
        highlighted_doc.paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
        highlighted_doc.save(highlighted)

        artifacts = [
            ("manuscript_clean", clean, "Manuscript_Clean.docx", None),
            ("manuscript_highlighted", highlighted, "Manuscript_Highlighted.docx", None),
        ]
        figure = sources / "Figure_1.png"
        figure.write_bytes(b"formal figure")
        artifacts.append(("main_figure", figure, "Main_Figures_and_Tables/Figure_1.png", None))
        table = sources / "Table_1.docx"
        self.make_docx(table, "Table 1")
        artifacts.append(("main_table", table, "Main_Figures_and_Tables/Table_1.docx", None))
        supplementary_figure = sources / "Supplementary_Figure_1.png"
        supplementary_figure.write_bytes(b"formal supplementary figure")
        artifacts.append(("supplementary_figure", supplementary_figure, "Supplementary_Materials/Supplementary_Figure_1.png", None))
        supplementary_table = sources / "Supplementary_Table_1.docx"
        self.make_docx(supplementary_table, "Supplementary Table 1")
        artifacts.append(("supplementary_table", supplementary_table, "Supplementary_Materials/Supplementary_Table_1.docx", None))
        for reviewer_number in (1, 2):
            response = sources / f"Response_to_Reviewer_{reviewer_number}.docx"
            self.make_response(response, reviewer_number)
            artifacts.append(
                (
                    "reviewer_response",
                    response,
                    f"Response_to_Reviewers/Response_to_Reviewer_{reviewer_number}.docx",
                    f"R{reviewer_number}",
                )
            )
        checklist = sources / "Checklist.docx"
        self.make_docx(checklist, "Completed submission checklist")
        artifacts.append(("submission_document", checklist, "Submission_Documents/Checklist.docx", None))

        policy = {
            "schema_version": "1.0",
            "workflow_type": "reviewer_revision",
            "package_name": "Submission_Package",
            "required_roles": ["manuscript_clean", "manuscript_highlighted", "reviewer_response"],
            "allowed_extensions": [".docx", ".png", ".tif", ".tiff", ".jpg", ".jpeg", ".pdf", ".xlsx"],
            "prohibited_extensions": [".json", ".md", ".csv", ".log", ".py", ".ps1", ".xml", ".tmp", ".bak"],
            "prohibited_name_patterns": [
                "^\\.",
                "~\\$",
                "(?i)(audit|ledger|work[-_ ]?log|execution|process|draft|backup|temp|cache)",
            ],
            "registered_artifacts": [],
        }
        for role, source, relative_path, reviewer_id in artifacts:
            item = {
                "role": role,
                "source": str(source),
                "relative_path": relative_path,
                "sha256": sha256_file(source),
            }
            if reviewer_id:
                item["reviewer_id"] = reviewer_id
            policy["registered_artifacts"].append(item)
        return policy, sources

    def test_safe_assembly_and_clean_validation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            policy, _ = self.make_project(root)
            package = assemble_package(policy, root / "delivery")
            self.assertEqual("Submission_Package", package.name)
            self.assertEqual([], validate_submission_package(package, policy))

    def test_contamination_extensions_and_names_are_rejected(self) -> None:
        prohibited = [
            "audit.json",
            "notes.md",
            "ledger.csv",
            "run.log",
            "helper.py",
            "scratch.tmp",
            "manuscript.bak",
            ".hidden.docx",
            "workspace.xml",
            "draft.docx",
        ]
        for filename in prohibited:
            with self.subTest(filename=filename), tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                policy, _ = self.make_project(root)
                package = assemble_package(policy, root / "delivery")
                (package / filename).write_text("internal", encoding="utf-8")
                codes = {item["code"] for item in validate_submission_package(package, policy)}
                self.assertTrue(codes & {"unknown_artifact", "prohibited_extension", "prohibited_name", "hidden_artifact"})

    def test_missing_reviewer_response_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            policy, _ = self.make_project(root)
            package = assemble_package(policy, root / "delivery")
            (package / "Response_to_Reviewers/Response_to_Reviewer_2.docx").unlink()
            codes = {item["code"] for item in validate_submission_package(package, policy)}
            self.assertIn("missing_registered_artifact", codes)
            self.assertIn("reviewer_response_mismatch", codes)

    def test_unresolved_placeholder_and_tracked_change_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            policy, _ = self.make_project(root)
            package = assemble_package(policy, root / "delivery")
            response = package / "Response_to_Reviewers/Response_to_Reviewer_1.docx"
            document = Document(response)
            document.add_paragraph("[[INSERT LOCATION]]")
            document.save(response)
            codes = {item["code"] for item in validate_submission_package(package, policy, verify_hashes=False)}
            self.assertIn("unresolved_placeholder", codes)

            manuscript = package / "Manuscript_Clean.docx"
            temporary_docx = manuscript.with_suffix(".tmp")
            with zipfile.ZipFile(manuscript, "r") as source, zipfile.ZipFile(temporary_docx, "w") as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename == "word/document.xml":
                        data = data.replace(b"<w:r>", b"<w:ins><w:r>", 1).replace(b"</w:r>", b"</w:r></w:ins>", 1)
                    target.writestr(item, data)
            temporary_docx.replace(manuscript)
            codes = {item["code"] for item in validate_submission_package(package, policy, verify_hashes=False)}
            self.assertIn("tracked_change_marker", codes)


if __name__ == "__main__":
    unittest.main()
