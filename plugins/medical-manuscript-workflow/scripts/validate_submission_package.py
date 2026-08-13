from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

from docx import Document
from lxml import etree

try:
    from .compare_manuscript_variants import compare_variants
    from .workflow_common import cli_result, issue, load_json, sha256_file
except ImportError:
    from compare_manuscript_variants import compare_variants
    from workflow_common import cli_result, issue, load_json, sha256_file


PLACEHOLDER_PATTERNS = [
    re.compile(r"\b(?:TODO|TBD|PLACEHOLDER)\b", re.IGNORECASE),
    re.compile(r"\[\[[^\]]+\]\]"),
    re.compile(r"\{\{[^}]+\}\}"),
]
TRACKED_TAGS = {"ins", "del", "moveFrom", "moveTo"}
FORMAL_ROOTS = {
    "Main_Figures_and_Tables",
    "Supplementary_Materials",
    "Response_to_Reviewers",
    "Submission_Documents",
}
ROLE_ROOTS = {
    "main_figure": "Main_Figures_and_Tables",
    "main_table": "Main_Figures_and_Tables",
    "supplementary_figure": "Supplementary_Materials",
    "supplementary_table": "Supplementary_Materials",
    "reviewer_response": "Response_to_Reviewers",
    "submission_document": "Submission_Documents",
}


def _safe_relative(value: str) -> bool:
    path = PurePosixPath(value.replace("\\", "/"))
    return not path.is_absolute() and bool(path.parts) and all(part not in {"", ".", ".."} for part in path.parts)


def _docx_text_and_markers(path: Path) -> tuple[str, set[str], bool, bool]:
    text = ""
    try:
        document = Document(path)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
    except Exception:
        return "", set(), False, False
    tracked: set[str] = set()
    has_hidden = False
    has_comments = False
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = set(archive.namelist())
            has_comments = "word/comments.xml" in names
            for name in names:
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                try:
                    root = etree.fromstring(archive.read(name))
                except etree.XMLSyntaxError:
                    continue
                for node in root.iter():
                    local = etree.QName(node).localname
                    if local in TRACKED_TAGS:
                        tracked.add(local)
                    elif local == "vanish":
                        has_hidden = True
                    elif local in {"commentRangeStart", "commentRangeEnd", "commentReference"}:
                        has_comments = True
    except zipfile.BadZipFile:
        pass
    return text, tracked, has_hidden, has_comments


def _response_structure_issues(path: Path, relative_path: str) -> list[dict[str, Any]]:
    try:
        document = Document(path)
    except Exception as exc:
        return [issue("invalid_docx", f"Unable to open formal DOCX: {exc}", relative_path)]
    headings = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip() in {"Reviewer Comment", "Response", "Location of Revision"}
    ]
    expected_unit = ["Reviewer Comment", "Response", "Location of Revision"]
    if not headings or len(headings) % 3 or any(headings[index : index + 3] != expected_unit for index in range(0, len(headings), 3)):
        return [issue("invalid_response_sequence", "Reviewer response must repeat Comment, Response, and Location headings in order.", relative_path)]
    return []


def validate_submission_package(
    package_dir: Path,
    policy: dict[str, Any],
    *,
    verify_hashes: bool = True,
) -> list[dict[str, Any]]:
    package_dir = Path(package_dir)
    issues: list[dict[str, Any]] = []
    if not package_dir.is_dir():
        return [issue("missing_package", "Submission package directory does not exist.", str(package_dir))]

    registrations: dict[str, dict[str, Any]] = {}
    for index, artifact in enumerate(policy.get("registered_artifacts", [])):
        relative = artifact.get("relative_path", "")
        if not isinstance(relative, str) or not _safe_relative(relative):
            issues.append(issue("unsafe_relative_path", "Registered relative path is unsafe.", f"registered_artifacts[{index}].relative_path"))
            continue
        normalized = PurePosixPath(relative.replace("\\", "/")).as_posix()
        if normalized in registrations:
            issues.append(issue("duplicate_relative_path", "Registered formal paths must be unique.", normalized))
        registrations[normalized] = artifact

    actual_files = {
        path.relative_to(package_dir).as_posix(): path
        for path in package_dir.rglob("*")
        if path.is_file()
    }
    expected = set(registrations)
    actual = set(actual_files)
    for relative in sorted(expected - actual):
        issues.append(issue("missing_registered_artifact", "Registered formal artifact is missing.", relative))
    for relative in sorted(actual - expected):
        issues.append(issue("unknown_artifact", "File is not present in the formal allowlist.", relative))

    allowed_extensions = {str(value).lower() for value in policy.get("allowed_extensions", [])}
    prohibited_extensions = {str(value).lower() for value in policy.get("prohibited_extensions", [])}
    compiled_patterns: list[re.Pattern[str]] = []
    for index, pattern in enumerate(policy.get("prohibited_name_patterns", [])):
        try:
            compiled_patterns.append(re.compile(pattern))
        except re.error as exc:
            issues.append(issue("invalid_policy_pattern", f"Invalid prohibited-name pattern: {exc}", f"prohibited_name_patterns[{index}]"))

    roles_present: dict[str, list[str]] = {}
    for relative, path in actual_files.items():
        parts = PurePosixPath(relative).parts
        if any(part.startswith(".") for part in parts):
            issues.append(issue("hidden_artifact", "Hidden files are not allowed in the submission package.", relative))
        extension = path.suffix.lower()
        if extension in prohibited_extensions:
            issues.append(issue("prohibited_extension", "File extension is prohibited in the submission package.", relative))
        if extension not in allowed_extensions:
            issues.append(issue("unapproved_extension", "File extension is not in the formal allowlist.", relative))
        if any(pattern.search(relative) for pattern in compiled_patterns):
            issues.append(issue("prohibited_name", "File name matches an internal or temporary artifact pattern.", relative))

        registration = registrations.get(relative)
        if registration:
            role = registration.get("role")
            roles_present.setdefault(role, []).append(relative)
            expected_root = ROLE_ROOTS.get(role)
            if expected_root and (not parts or parts[0] != expected_root):
                issues.append(issue("invalid_role_location", f"Artifact role {role} must be stored under {expected_root}.", relative))
            if role == "manuscript_clean" and relative != "Manuscript_Clean.docx":
                issues.append(issue("invalid_manuscript_name", "Clean manuscript must use the approved formal filename.", relative))
            if role == "manuscript_highlighted" and relative != "Manuscript_Highlighted.docx":
                issues.append(issue("invalid_manuscript_name", "Highlighted manuscript must use the approved formal filename.", relative))
            if verify_hashes and sha256_file(path).lower() != str(registration.get("sha256", "")).lower():
                issues.append(issue("artifact_hash_mismatch", "Formal artifact differs from its registered release hash.", relative))

        if extension == ".docx":
            text, tracked, hidden, comments = _docx_text_and_markers(path)
            if any(pattern.search(text) for pattern in PLACEHOLDER_PATTERNS):
                issues.append(issue("unresolved_placeholder", "Formal DOCX contains an unresolved placeholder.", relative))
            if tracked:
                issues.append(issue("tracked_change_marker", "Formal DOCX contains unresolved tracked-change markup.", relative, markers=sorted(tracked)))
            if hidden:
                issues.append(issue("hidden_text", "Formal DOCX contains hidden text markup.", relative))
            if comments:
                issues.append(issue("word_comments_present", "Formal DOCX contains Word comments.", relative))
            if registration and registration.get("role") == "reviewer_response":
                issues.extend(_response_structure_issues(path, relative))

    for required_role in policy.get("required_roles", []):
        if not roles_present.get(required_role):
            issues.append(issue("missing_required_role", f"Required formal artifact role is missing: {required_role}.", "registered_artifacts"))

    if policy.get("workflow_type") == "reviewer_revision":
        clean_path = package_dir / "Manuscript_Clean.docx"
        highlighted_path = package_dir / "Manuscript_Highlighted.docx"
        if clean_path.exists() and highlighted_path.exists():
            for variant_issue in compare_variants(clean_path, highlighted_path):
                variant_issue = dict(variant_issue)
                variant_issue["path"] = "Manuscript_Clean.docx | Manuscript_Highlighted.docx"
                issues.append(variant_issue)
        expected_reviewers = {
            artifact.get("reviewer_id")
            for artifact in registrations.values()
            if artifact.get("role") == "reviewer_response" and artifact.get("reviewer_id")
        }
        actual_responses = roles_present.get("reviewer_response", [])
        actual_reviewer_ids = {
            registrations[path].get("reviewer_id")
            for path in actual_responses
            if path in registrations and registrations[path].get("reviewer_id")
        }
        if len(actual_responses) != len(expected_reviewers) or actual_reviewer_ids != expected_reviewers:
            issues.append(issue("reviewer_response_mismatch", "Each registered Reviewer must have exactly one formal response DOCX.", "Response_to_Reviewers"))
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate a pure formal submission package.")
    parser.add_argument("package_dir")
    parser.add_argument("policy")
    parser.add_argument("--skip-hashes", action="store_true")
    args = parser.parse_args()
    return cli_result(
        validate_submission_package(
            Path(args.package_dir),
            load_json(args.policy),
            verify_hashes=not args.skip_hashes,
        )
    )


if __name__ == "__main__":
    raise SystemExit(main())
