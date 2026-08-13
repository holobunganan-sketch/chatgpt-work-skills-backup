from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

try:
    from .validate_revision_release import release_issues
    from .workflow_common import load_json
except ImportError:
    from validate_revision_release import release_issues
    from workflow_common import load_json


def _location_text(location: dict[str, Any]) -> str:
    if location.get("status") == "not_applicable":
        return "No change to the manuscript."
    rendered: list[str] = []
    for entry in location.get("entries", []):
        parts = [entry.get("artifact", ""), entry.get("section", "")]
        if entry.get("page") is not None:
            parts.append(f"Page {entry['page']}")
        if entry.get("paragraph") is not None:
            parts.append(f"Paragraph {entry['paragraph']}")
        parts.append(entry.get("anchor", ""))
        rendered.append("; ".join(part for part in parts if part))
    return "\n".join(rendered)


def _set_document_defaults(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"


def build_responses(ledger: dict[str, Any], output_dir: Path) -> list[Path]:
    issues = release_issues(ledger)
    if issues:
        codes = ", ".join(sorted({item["code"] for item in issues}))
        raise ValueError(f"Reviewer ledger is not release-ready: {codes}")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_paths: list[Path] = []
    reviewers = sorted(ledger["reviewers"], key=lambda item: item["sequence"])
    for reviewer in reviewers:
        sequence = reviewer["sequence"]
        destination = output_dir / f"Response_to_Reviewer_{sequence}.docx"
        if destination.exists():
            raise FileExistsError(f"Refusing to overwrite existing response file: {destination}")

        document = Document()
        _set_document_defaults(document)
        document.core_properties.title = f"Response to Reviewer {sequence}"
        document.add_heading(f"Response to Reviewer {sequence}", level=1)
        for comment in sorted(reviewer["comments"], key=lambda item: item["order"]):
            document.add_heading("Reviewer Comment", level=2)
            document.add_paragraph(comment["verbatim_comment"])
            document.add_heading("Response", level=2)
            document.add_paragraph(comment["response"])
            document.add_heading("Location of Revision", level=2)
            document.add_paragraph(_location_text(comment["location"]))
        document.save(destination)
        output_paths.append(destination)
    return output_paths


def main() -> int:
    parser = argparse.ArgumentParser(description="Build one formal response DOCX per Reviewer.")
    parser.add_argument("ledger", help="Release-ready reviewer ledger JSON")
    parser.add_argument("output_dir", help="Destination directory for formal response DOCX files")
    args = parser.parse_args()
    for path in build_responses(load_json(args.ledger), Path(args.output_dir)):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
