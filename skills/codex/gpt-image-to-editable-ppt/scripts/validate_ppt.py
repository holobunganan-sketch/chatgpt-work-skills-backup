from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from common import emu_to_inches, load_data, occupancy_ratio, rect_intersection


def get_font_sizes(shape):
    sizes = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for p in shape.text_frame.paragraphs:
        for run in p.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    return sizes


def get_font_colors(shape):
    colors = []
    if not getattr(shape, "has_text_frame", False):
        return colors
    for p in shape.text_frame.paragraphs:
        for run in p.runs:
            try:
                if run.font.color.type and run.font.color.rgb:
                    colors.append(f"#{run.font.color.rgb}")
            except Exception:
                pass
    return colors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", required=True)
    parser.add_argument("--design", required=True)
    parser.add_argument("--spec")
    parser.add_argument("--report")
    args = parser.parse_args()

    design = load_data(args.design)
    spec = load_data(args.spec) if args.spec else {}
    prs = Presentation(args.pptx)
    sw, sh = emu_to_inches(prs.slide_width), emu_to_inches(prs.slide_height)
    min_font = float(design["typography"].get("minimum_font_size_pt", 14))
    default_text = design["typography"].get("default_text_color", "#000000").upper()
    density = design.get("layout_density", {})
    min_occ = float(density.get("minimum_occupancy", 0.50))
    max_occ = float(density.get("maximum_occupancy", 0.86))
    errors, warnings = [], []

    if spec and len(prs.slides) != len(spec.get("slides", [])):
        errors.append({"code": "slide_count_mismatch", "pptx": len(prs.slides), "spec": len(spec.get("slides", []))})

    for idx, slide in enumerate(prs.slides, start=1):
        slide_spec = spec.get("slides", [{}] * len(prs.slides))[idx - 1] if spec else {}
        page_type = slide_spec.get("page_type")
        rects = []
        native_text = 0
        full_slide_images = 0
        named = {shape.name for shape in slide.shapes}
        shape_records = []

        for shape in slide.shapes:
            x, y, w, h = map(emu_to_inches, (shape.left, shape.top, shape.width, shape.height))
            if x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + h > sh + 0.01:
                errors.append({"code": "out_of_bounds", "slide": idx, "shape": shape.name, "box": [round(x,3), round(y,3), round(w,3), round(h,3)]})
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and w * h >= sw * sh * 0.92:
                full_slide_images += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                native_text += 1
                sizes = get_font_sizes(shape)
                if sizes and min(sizes) < min_font - 0.01:
                    errors.append({"code": "font_too_small", "slide": idx, "shape": shape.name, "minimum_found": round(min(sizes), 2), "required": min_font})
                colors = get_font_colors(shape)
                if default_text == "#000000" and any(c.upper() != default_text for c in colors):
                    warnings.append({"code": "non_default_text_color", "slide": idx, "shape": shape.name, "colors": sorted(set(colors))})
                # Approximate overflow risk; PowerPoint rendering remains authoritative.
                fs = min(sizes) if sizes else design["typography"].get("body_font_size_pt", 18)
                capacity = max(1, (w * 72 / (fs * 0.9)) * (h * 72 / (fs * 1.5)))
                if len(shape.text) > capacity * 1.8:
                    warnings.append({"code": "possible_text_overflow", "slide": idx, "shape": shape.name, "characters": len(shape.text), "capacity_estimate": int(capacity)})
            if shape.name not in {"PAGE_TITLE", "PAGE_DIVIDER"} and shape.shape_type != MSO_SHAPE_TYPE.LINE:
                rects.append((x, y, w, h))
            allow_overlap = "|overlap" in shape.name
            shape_records.append((shape, (x, y, x+w, y+h), max(1e-9, w*h), allow_overlap))

        if page_type == "content" and not slide_spec.get("suppress_standard_chrome", False):
            if "PAGE_TITLE" not in named:
                errors.append({"code": "missing_page_title", "slide": idx})
            if "PAGE_DIVIDER" not in named:
                errors.append({"code": "missing_page_divider", "slide": idx})

        if full_slide_images and spec.get("deck", {}).get("delivery_mode", "hybrid") != "image":
            errors.append({"code": "full_slide_raster_forbidden", "slide": idx})
        if full_slide_images and native_text == 0:
            warnings.append({"code": "low_editability", "slide": idx})

        occ = occupancy_ratio(rects, sw, sh)
        if page_type == "content" and occ < min_occ:
            warnings.append({"code": "low_occupancy", "slide": idx, "value": round(occ, 3)})
        if occ > max_occ:
            warnings.append({"code": "high_occupancy", "slide": idx, "value": round(occ, 3)})

        for i, (a, ar, aa, a_allow) in enumerate(shape_records):
            if a_allow or a.name in {"PAGE_TITLE", "PAGE_DIVIDER"} or a.shape_type == MSO_SHAPE_TYPE.LINE:
                continue
            for b, br, ba, b_allow in shape_records[i+1:]:
                if b_allow or b.name in {"PAGE_TITLE", "PAGE_DIVIDER"} or b.shape_type == MSO_SHAPE_TYPE.LINE:
                    continue
                overlap = rect_intersection(ar, br)
                if overlap / min(aa, ba) > 0.35:
                    warnings.append({"code": "possible_overlap", "slide": idx, "shapes": [a.name, b.name], "ratio": round(overlap/min(aa,ba),3)})

    report = {"valid": not errors, "errors": errors, "warnings": warnings, "slide_count": len(prs.slides)}
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        Path(args.report).write_text(text, encoding="utf-8")
    print(text)
    return 0 if not errors else 2


if __name__ == "__main__":
    raise SystemExit(main())
