from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
from typing import Any, Dict

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from common import hex_to_rgb, load_data

SHAPE_MAP = {
    "rect": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
    "rectangle": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
    "round_rect": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    "rounded_rectangle": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_AUTO_SHAPE_TYPE.OVAL,
    "circle": MSO_AUTO_SHAPE_TYPE.OVAL,
    "chevron": MSO_AUTO_SHAPE_TYPE.CHEVRON,
    "hexagon": MSO_AUTO_SHAPE_TYPE.HEXAGON,
    "triangle": MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_AUTO_SHAPE_TYPE.DIAMOND,
}

CHART_MAP = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor(*hex_to_rgb(hex_value))


def set_run_font(run, family: str, size_pt: float, color: str, bold: bool = False):
    run.font.name = family
    run.font.size = Pt(size_pt)
    run.font.bold = bool(bold)
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set(qn("a:lang"), "zh-CN")
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            from pptx.oxml.xmlchemy import OxmlElement
            node = OxmlElement(tag)
            rpr.append(node)
        node.set("typeface", family)


def style_text_frame(tf, text: str, style: Dict[str, Any], design: Dict[str, Any]):
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(float(style.get("margin_left_in", 0.08)))
    tf.margin_right = Inches(float(style.get("margin_right_in", 0.08)))
    tf.margin_top = Inches(float(style.get("margin_top_in", 0.04)))
    tf.margin_bottom = Inches(float(style.get("margin_bottom_in", 0.04)))
    valign = style.get("valign", "middle")
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM
    }.get(valign, MSO_ANCHOR.MIDDLE)
    tf.text = text or ""
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(style.get("align", "left"), PP_ALIGN.LEFT)
    line_spacing = float(style.get("line_spacing", design["typography"].get("preferred_line_spacing", 1.5)))
    family = style.get("font_family", design["typography"]["font_family"])
    size = max(float(style.get("font_size_pt", design["typography"].get("body_font_size_pt", 18))), float(design["typography"].get("minimum_font_size_pt", 14)))
    color = style.get("color", design["typography"].get("default_text_color", "#000000"))
    bold = style.get("bold", False)
    for p in tf.paragraphs:
        p.alignment = alignment
        p.line_spacing = line_spacing
        if not p.runs:
            p.add_run()
        for run in p.runs:
            set_run_font(run, family, size, color, bold)


def clear_slides(prs: Presentation):
    slide_ids = prs.slides._sldIdLst
    for sld_id in list(slide_ids):
        rel_id = sld_id.rId
        prs.part.drop_rel(rel_id)
        slide_ids.remove(sld_id)


def choose_layout(prs: Presentation, page_type: str):
    keywords = {
        "cover": ["title slide", "cover", "标题"],
        "section": ["section", "章节"],
        "agenda": ["title and content", "目录", "agenda"],
        "content": ["blank", "空白", "title only", "仅标题"],
        "closing": ["title slide", "closing", "结束"],
    }.get(page_type, ["blank"])
    for key in keywords:
        for layout in prs.slide_layouts:
            if key.lower() in (layout.name or "").lower():
                return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


def set_slide_background(slide, color: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_standard_chrome(slide, slide_spec, design):
    layout = design["standard_content_layout"]
    title = slide.shapes.add_textbox(Inches(layout["title_x_in"]), Inches(layout["title_y_in"]), Inches(layout["title_w_in"]), Inches(layout["title_h_in"]))
    title.name = "PAGE_TITLE"
    style_text_frame(title.text_frame, slide_spec.get("title", ""), {
        "font_size_pt": design["typography"].get("title_font_size_pt", 28),
        "font_family": design["typography"]["font_family"],
        "color": design["typography"].get("default_text_color", "#000000"),
        "bold": design["typography"].get("title_bold", True),
        "align": "left", "valign": "middle", "line_spacing": 1.0,
    }, design)
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(layout["divider_x_in"]), Inches(layout["divider_y_in"]),
        Inches(layout["divider_x_in"] + layout["divider_w_in"]), Inches(layout["divider_y_in"]),
    )
    line.name = "PAGE_DIVIDER"
    line.line.color.rgb = rgb(layout.get("divider_color", design["palette"]["primary"]))
    line.line.width = Pt(layout.get("divider_width_pt", 1.5))


def add_text(slide, element, design):
    shape = slide.shapes.add_textbox(Inches(element["x"]), Inches(element["y"]), Inches(element["w"]), Inches(element["h"]))
    shape.name = f"EL|{element['id']}|text|{'overlap' if element.get('allow_overlap') else 'normal'}"
    style_text_frame(shape.text_frame, element.get("text", ""), element, design)
    return shape


def add_shape(slide, element, design):
    shape_type = SHAPE_MAP.get(element.get("shape", "round_rect"), MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
    shape = slide.shapes.add_shape(shape_type, Inches(element["x"]), Inches(element["y"]), Inches(element["w"]), Inches(element["h"]))
    shape.name = f"EL|{element['id']}|shape|{'overlap' if element.get('allow_overlap') else 'normal'}"
    fill_color = element.get("fill", design["palette"].get("neutral_fill", "#F2F4F7"))
    if fill_color == "transparent":
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill_color)
    shape.line.color.rgb = rgb(element.get("line_color", fill_color if fill_color != "transparent" else design["palette"].get("neutral_line", "#B8C0CC")))
    shape.line.width = Pt(float(element.get("line_width_pt", 1.0)))
    if element.get("text") is not None:
        style_text_frame(shape.text_frame, element.get("text", ""), element, design)
    return shape


def add_line(slide, element, design):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(element["x"]), Inches(element["y"]),
        Inches(element["x"] + element["w"]), Inches(element["y"] + element.get("h", 0)),
    )
    line.name = f"EL|{element['id']}|line|overlap"
    line.line.color.rgb = rgb(element.get("line_color", design["palette"]["primary"]))
    line.line.width = Pt(float(element.get("line_width_pt", 1.5)))
    return line


def add_image(slide, element, spec_dir: Path):
    path = Path(element["path"])
    if not path.is_absolute():
        path = spec_dir / path
    if not path.exists():
        raise FileNotFoundError(f"Image asset not found: {path}")
    pic = slide.shapes.add_picture(str(path), Inches(element["x"]), Inches(element["y"]), Inches(element["w"]), Inches(element["h"]))
    pic.name = f"EL|{element['id']}|image|{'overlap' if element.get('allow_overlap') else 'normal'}"
    return pic


def add_table(slide, element, design):
    data = element.get("data", [])
    if not data or not isinstance(data, list):
        raise ValueError(f"Table {element['id']} has no row data")
    rows = len(data)
    cols = max(len(row) for row in data)
    shape = slide.shapes.add_table(rows, cols, Inches(element["x"]), Inches(element["y"]), Inches(element["w"]), Inches(element["h"]))
    shape.name = f"EL|{element['id']}|table|normal"
    table = shape.table
    for r in range(rows):
        for c in range(cols):
            value = str(data[r][c]) if c < len(data[r]) else ""
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(element.get("header_fill", design["palette"]["primary_light"]) if r == 0 else element.get("fill", "#FFFFFF"))
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER if element.get("align", "center") == "center" else PP_ALIGN.LEFT
                p.line_spacing = float(element.get("line_spacing", design["typography"].get("preferred_line_spacing", 1.5)))
                if p.runs:
                    set_run_font(p.runs[0], element.get("font_family", design["typography"]["font_family"]), max(float(element.get("font_size_pt", 14)), float(design["typography"].get("minimum_font_size_pt", 14))), element.get("color", design["typography"].get("default_text_color", "#000000")), r == 0)
    return shape


def add_chart(slide, element, design):
    payload = element.get("data", {})
    categories = payload.get("categories", [])
    series = payload.get("series", [])
    if not categories or not series:
        raise ValueError(f"Chart {element['id']} requires categories and series")
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for item in series:
        chart_data.add_series(item.get("name", "Series"), item.get("values", []))
    chart_type = CHART_MAP.get(element.get("chart_type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
    chart = slide.shapes.add_chart(chart_type, Inches(element["x"]), Inches(element["y"]), Inches(element["w"]), Inches(element["h"]), chart_data).chart
    chart.has_legend = bool(element.get("show_legend", len(series) > 1))
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.has_title = bool(element.get("title"))
    if chart.has_title:
        chart.chart_title.text_frame.text = element["title"]
    shape = slide.shapes[-1]
    shape.name = f"EL|{element['id']}|chart|normal"
    return shape


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--spec", required=True)
    parser.add_argument("--design", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--template")
    parser.add_argument("--keep-template-slides", action="store_true")
    args = parser.parse_args()

    spec_path = Path(args.spec).resolve()
    spec = load_data(spec_path)
    design = load_data(args.design)
    template = args.template or spec.get("deck", {}).get("template_path")
    if template:
        prs = Presentation(str(Path(template).resolve()))
        if not args.keep_template_slides:
            clear_slides(prs)
    else:
        prs = Presentation()
        clear_slides(prs)

    prs.slide_width = Inches(float(design["slide"].get("width_in", 13.333333)))
    prs.slide_height = Inches(float(design["slide"].get("height_in", 7.5)))
    background = design["slide"].get("background_color", "#FFFFFF")

    for slide_spec in spec.get("slides", []):
        page_type = slide_spec.get("page_type", "content")
        slide = prs.slides.add_slide(choose_layout(prs, page_type))
        # Clear editable placeholders created by the selected layout, while retaining master graphics.
        for shape in list(slide.shapes):
            if getattr(shape, "is_placeholder", False):
                sp = shape._element
                sp.getparent().remove(sp)
        set_slide_background(slide, slide_spec.get("background_color", background))
        if page_type == "content" and not slide_spec.get("suppress_standard_chrome", False):
            add_standard_chrome(slide, slide_spec, design)

        for element in sorted(slide_spec.get("elements", []), key=lambda e: e.get("z", 0)):
            etype = element.get("type")
            if etype == "text":
                add_text(slide, element, design)
            elif etype == "shape":
                add_shape(slide, element, design)
            elif etype == "line":
                add_line(slide, element, design)
            elif etype == "image":
                add_image(slide, element, spec_path.parent)
            elif etype == "table":
                add_table(slide, element, design)
            elif etype == "chart":
                add_chart(slide, element, design)
            else:
                raise ValueError(f"Unsupported element type: {etype}")

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(json.dumps({"output": str(out.resolve()), "slides": len(prs.slides)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
