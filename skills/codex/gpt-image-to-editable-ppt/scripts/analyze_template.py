from __future__ import annotations

import argparse
import json
import statistics
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation

from common import emu_to_inches, save_data

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def extract_theme(pptx_path: Path) -> dict:
    result = {"colors": {}, "fonts": {}}
    try:
        with zipfile.ZipFile(pptx_path) as zf:
            theme_names = sorted(n for n in zf.namelist() if n.startswith("ppt/theme/theme") and n.endswith(".xml"))
            if not theme_names:
                return result
            root = ET.fromstring(zf.read(theme_names[0]))
            clr_scheme = root.find(".//a:clrScheme", NS)
            if clr_scheme is not None:
                for child in list(clr_scheme):
                    name = child.tag.split("}")[-1]
                    color_node = next(iter(child), None)
                    if color_node is not None:
                        value = color_node.attrib.get("val") or color_node.attrib.get("lastClr")
                        if value:
                            result["colors"][name] = f"#{value.upper()}"
            major = root.find(".//a:fontScheme/a:majorFont", NS)
            minor = root.find(".//a:fontScheme/a:minorFont", NS)
            for label, node in [("major", major), ("minor", minor)]:
                if node is None:
                    continue
                latin = node.find("a:latin", NS)
                ea = node.find("a:ea", NS)
                result["fonts"][label] = {
                    "latin": latin.attrib.get("typeface", "") if latin is not None else "",
                    "east_asian": ea.attrib.get("typeface", "") if ea is not None else "",
                }
    except Exception as exc:
        result["error"] = str(exc)
    return result


def median_box(boxes):
    if not boxes:
        return None
    return {
        "x_in": round(statistics.median(b[0] for b in boxes), 3),
        "y_in": round(statistics.median(b[1] for b in boxes), 3),
        "w_in": round(statistics.median(b[2] for b in boxes), 3),
        "h_in": round(statistics.median(b[3] for b in boxes), 3),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    path = Path(args.template)
    prs = Presentation(str(path))
    title_boxes = []
    body_boxes = []
    shape_signatures = Counter()
    background_colors = Counter()
    sample_slides = []

    for idx, slide in enumerate(prs.slides):
        slide_summary = {"index": idx + 1, "layout": slide.slide_layout.name, "shapes": []}
        try:
            fill = slide.background.fill
            if fill.type and getattr(fill.fore_color, "rgb", None):
                background_colors[f"#{fill.fore_color.rgb}"] += 1
        except Exception:
            pass
        for shape in slide.shapes:
            box = tuple(round(emu_to_inches(v), 3) for v in (shape.left, shape.top, shape.width, shape.height))
            text = shape.text.strip() if getattr(shape, "has_text_frame", False) else ""
            ph_type = None
            if getattr(shape, "is_placeholder", False):
                try:
                    ph_type = str(shape.placeholder_format.type)
                except Exception:
                    pass
            if ph_type and "TITLE" in ph_type.upper():
                title_boxes.append(box)
            elif text and len(text) > 20:
                body_boxes.append(box)
            sig = (shape.shape_type, round(box[0], 1), round(box[1], 1), round(box[2], 1), round(box[3], 1))
            shape_signatures[sig] += 1
            slide_summary["shapes"].append({"name": shape.name, "type": str(shape.shape_type), "box": box, "text_preview": text[:80]})
        if idx < 8:
            sample_slides.append(slide_summary)

    recurring = []
    threshold = max(2, int(len(prs.slides) * 0.5)) if len(prs.slides) else 2
    for sig, count in shape_signatures.most_common(30):
        if count >= threshold:
            recurring.append({"signature": list(sig), "count": count})

    theme = extract_theme(path)
    theme_font = theme.get("fonts", {}).get("minor", {}).get("east_asian") or theme.get("fonts", {}).get("major", {}).get("east_asian")
    theme_text = theme.get("colors", {}).get("dk1", "#000000")
    theme_bg = theme.get("colors", {}).get("lt1", "#FFFFFF")

    profile = {
        "source_template": str(path.resolve()),
        "slide_count": len(prs.slides),
        "slide_size": {
            "width_in": round(emu_to_inches(prs.slide_width), 4),
            "height_in": round(emu_to_inches(prs.slide_height), 4),
        },
        "layouts": [{"index": i, "name": layout.name} for i, layout in enumerate(prs.slide_layouts)],
        "theme": theme,
        "representative_title_box": median_box(title_boxes),
        "representative_body_box": median_box(body_boxes),
        "background_colors": background_colors.most_common(8),
        "recurring_shapes": recurring,
        "sample_slides": sample_slides,
        "design_system": {
            "slide": {
                "width_in": round(emu_to_inches(prs.slide_width), 4),
                "height_in": round(emu_to_inches(prs.slide_height), 4),
                "background_color": theme_bg,
            },
            "typography": {
                "font_family": theme_font or "Microsoft YaHei",
                "default_text_color": theme_text,
            },
            "palette": {
                "background": theme_bg,
                "text": theme_text,
                "primary": theme.get("colors", {}).get("accent1", "#2F6B9A"),
                "secondary": theme.get("colors", {}).get("accent2", "#4F8A7B"),
                "accent": theme.get("colors", {}).get("accent3", "#D29A42"),
            },
        },
    }
    if profile["representative_title_box"]:
        box = profile["representative_title_box"]
        profile["design_system"]["standard_content_layout"] = {
            "title_x_in": box["x_in"], "title_y_in": box["y_in"], "title_w_in": box["w_in"], "title_h_in": box["h_in"]
        }

    save_data(args.out, profile)
    print(json.dumps({"template": str(path), "out": args.out, "slides": len(prs.slides)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
