from __future__ import annotations

import json
import shutil
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from audit_delivery_package import audit as audit_delivery
from audit_figure_assets import audit as audit_figures
from audit_manuscript import audit as audit_manuscript
from enforce_docx_format import DEFAULT_PROFILE, enforce


def add_zotero_field(paragraph, item_count: int = 2) -> None:
    payload = {
        "citationID": "SELFTEST",
        "properties": {
            "formattedCitation": "[1,2]",
            "plainCitation": "[1,2]",
            "noteIndex": 0,
        },
        "citationItems": [
            {
                "id": index + 1,
                "uris": [f"http://zotero.org/users/local/items/TEST{index + 1}"],
            }
            for index in range(item_count)
        ],
    }
    instruction = " ADDIN ZOTERO_ITEM CSL_CITATION " + json.dumps(
        payload, ensure_ascii=False, separators=(",", ":")
    )

    run = paragraph.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run.append(begin)

    run = paragraph.add_run()._r
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run.append(instr)

    run = paragraph.add_run()._r
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run.append(separate)

    paragraph.add_run("[1,2]")

    run = paragraph.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)


def make_figure(path: Path, dpi: int = 300) -> None:
    image = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((220, 200, 1280, 700), outline="#1F4E79", width=12)
    draw.line((300, 620, 650, 390, 1050, 500, 1220, 270), fill="#C55A11", width=16)
    image.save(path, dpi=(dpi, dpi))


def make_source_docx(
    path: Path, figure_path: Path, citation_count: int = 2
) -> None:
    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    title = document.add_paragraph("示例医学研究 Manuscript")
    title.style = document.styles["Title"]
    document.add_heading("Introduction", level=1)
    document.add_paragraph("该领域仍存在需要回答的关键临床问题。")
    document.add_paragraph("既有证据为研究假设提供了基础。")
    document.add_paragraph("本研究旨在评价预设临床结局并明确其应用价值。")

    document.add_heading("Methods", level=1)
    document.add_heading("Study design and population", level=2)
    body = document.add_paragraph(
        "本研究计划采用回顾性队列设计并纳入120例患者，通过预设分析评价临床结局。"
    )
    for run in body.runs:
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(31, 78, 121)

    document.add_heading("Results", level=1)
    document.add_heading("Clinical outcomes", level=2)
    document.add_paragraph(
        "主要结局达到预设统计学阈值，结果见 Table 1 and Figure 1。"
    )

    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for column, value in enumerate(["Variable", "Estimate", "P value"]):
        table.cell(0, column).text = value
    for column, value in enumerate(["Age", "1.25", "0.035"]):
        table.cell(1, column).text = value
    for column, value in enumerate(["Sex", "0.98", "0.214"]):
        table.cell(2, column).text = value
    document.add_paragraph("Table 1 Baseline characteristics")

    document.add_paragraph("Figure 1 Clinical outcome distribution")
    document.add_picture(str(figure_path), width=Inches(5.5))

    document.add_heading("Discussion", level=1)
    document.add_paragraph(
        "本研究围绕预设临床问题采用回顾性队列方法，获得了主要结局结果，并为后续研究提供了整体证据。"
    )
    document.add_paragraph("主要结果可结合研究设计和目标人群进行解释。")
    document.add_paragraph("相关观察结果与临床路径具有潜在联系。")
    document.add_paragraph("研究结果可为后续验证性工作提供依据。")
    document.add_paragraph(
        "本研究存在两点局限：第一，单中心数据限制了外部推广；第二，残余混杂仍可能存在。"
    )

    document.add_heading("Conclusion", level=1)
    document.add_paragraph(
        "本研究为该领域建立了可复核的证据基础，未来可支持更广泛人群中的前瞻性验证与临床转化。"
    )

    document.add_heading("References", level=1)
    citation = document.add_paragraph("Dynamic citation: ")
    add_zotero_field(citation, item_count=citation_count)
    document.save(path)


def make_invalid_structure_docx(path: Path) -> None:
    document = Document()
    title = document.add_paragraph("结构边界反例")
    title.style = document.styles["Title"]
    document.add_heading("Introduction", level=1)
    document.add_paragraph("现有证据仍存在空白。")
    document.add_heading("Background detail", level=2)
    document.add_paragraph("本研究拟回答预设问题。")

    document.add_heading("Methods", level=1)
    document.add_heading("Study population", level=2)
    document.add_paragraph("本研究共计纳入120例患者。")
    document.add_paragraph("采用回归模型分析主要结局。")

    document.add_heading("Results", level=1)
    document.add_heading("Primary outcome", level=2)
    repeated_result = (
        "主要结局在两组之间存在具有统计学意义且达到预设阈值的差异。"
    )
    document.add_paragraph(repeated_result)
    document.add_paragraph("这些结果提示干预可能通过改善炎症机制产生作用。")

    document.add_heading("Discussion", level=1)
    document.add_paragraph("本研究完成了预设分析并获得主要结果。")
    document.add_heading("Interpretation detail", level=2)
    document.add_paragraph("结果具有潜在临床意义。")
    document.add_paragraph("相关机制仍需进一步研究。")
    document.add_paragraph("后续研究可扩大样本并延长随访。")

    document.add_heading("Conclusion", level=1)
    document.add_paragraph(repeated_result)
    document.add_paragraph("未来仍需开展验证性研究。")

    document.add_heading("References", level=1)
    citation = document.add_paragraph("Dynamic citation: ")
    add_zotero_field(citation)
    document.save(path)


def make_simple_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def margins_match(path: Path, expected: dict[str, float]) -> bool:
    document = Document(path)
    for section in document.sections:
        for side in ("top", "bottom", "left", "right"):
            margin = getattr(section, f"{side}_margin")
            if margin is None or abs(margin.cm - expected[side]) > 0.01:
                return False
    return True


def run_self_test() -> dict:
    temp_root = Path(tempfile.mkdtemp(prefix="format-manuscript-selftest-"))
    try:
        source = temp_root / "source.docx"
        formatted = temp_root / "formatted.docx"
        figure = temp_root / "Figure 1.png"
        make_figure(figure)
        make_source_docx(source, figure)

        unformatted_margin_report = audit_manuscript(
            source,
            DEFAULT_PROFILE,
            stage="draft",
            compare_source=None,
        )
        unformatted_margin_codes = {
            item["code"] for item in unformatted_margin_report["issues"]
        }
        enforce_report = enforce(source, formatted, DEFAULT_PROFILE)
        manuscript_report = audit_manuscript(
            formatted,
            DEFAULT_PROFILE,
            stage="final",
            compare_source=source,
        )
        figure_report = audit_figures(temp_root, 300)

        delivery = temp_root / "delivery"
        (delivery / "Figures").mkdir(parents=True)
        (delivery / "Tables").mkdir(parents=True)
        shutil.copy2(formatted, delivery / "Manuscript.docx")
        shutil.copy2(figure, delivery / "Figures" / "Figure 1.png")
        shutil.copy2(formatted, delivery / "Tables" / "Table 1.docx")
        make_simple_docx(
            delivery / "Submission Checklist.docx", "TRIPOD Checklist"
        )
        manifest = temp_root / "delivery-manifest.json"
        manifest.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "manuscript": "Manuscript.docx",
                    "artifacts": [
                        {
                            "role": "figure",
                            "label": "Figure 1",
                            "path": "Figures/Figure 1.png",
                        },
                        {
                            "role": "table",
                            "label": "Table 1",
                            "path": "Tables/Table 1.docx",
                        },
                        {
                            "role": "submission_checklist",
                            "label": "Submission Checklist",
                            "path": "Submission Checklist.docx",
                        },
                    ],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        delivery_report = audit_delivery(delivery, manifest)

        excessive_source = temp_root / "source-five-citations.docx"
        excessive_formatted = temp_root / "formatted-five-citations.docx"
        make_source_docx(excessive_source, figure, citation_count=5)
        enforce(excessive_source, excessive_formatted, DEFAULT_PROFILE)
        excessive_citation_report = audit_manuscript(
            excessive_formatted,
            DEFAULT_PROFILE,
            stage="final",
            compare_source=excessive_source,
        )

        low_dpi_dir = temp_root / "low-dpi"
        low_dpi_dir.mkdir()
        make_figure(low_dpi_dir / "Figure 2.png", dpi=96)
        low_dpi_report = audit_figures(low_dpi_dir, 300)

        invalid_source = temp_root / "invalid-structure-source.docx"
        invalid_formatted = temp_root / "invalid-structure-formatted.docx"
        make_invalid_structure_docx(invalid_source)
        enforce(invalid_source, invalid_formatted, DEFAULT_PROFILE)
        invalid_structure_report = audit_manuscript(
            invalid_formatted,
            DEFAULT_PROFILE,
            stage="final",
            compare_source=invalid_source,
        )
        invalid_codes = {
            item["code"] for item in invalid_structure_report["issues"]
        }
        required_structure_codes = {
            "INTRODUCTION_PARAGRAPH_COUNT",
            "INTRODUCTION_SUBHEADING_FORBIDDEN",
            "METHODS_SUBHEADING_PARAGRAPH_COUNT",
            "METHODS_POST_EXECUTION_RESULT",
            "RESULTS_SUBHEADING_PARAGRAPH_COUNT",
            "RESULTS_INTERPRETATION",
            "DISCUSSION_PARAGRAPH_COUNT",
            "DISCUSSION_SUBHEADING_FORBIDDEN",
            "DISCUSSION_LIMITATIONS_MISSING",
            "CONCLUSION_PARAGRAPH_COUNT",
            "CONCLUSION_REPEATS_PRIOR_TEXT",
        }

        override_profile = temp_root / "explicit-user-format-profile.json"
        override_payload = json.loads(DEFAULT_PROFILE.read_text(encoding="utf-8"))
        override_payload["structure"]["enabled"] = False
        override_margins = {
            "top": 2.0,
            "bottom": 2.0,
            "left": 2.5,
            "right": 2.5,
        }
        override_payload["document"]["page_margins_cm"] = override_margins
        override_profile.write_text(
            json.dumps(override_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        override_formatted = temp_root / "explicit-user-format.docx"
        enforce(invalid_source, override_formatted, override_profile)
        override_report = audit_manuscript(
            override_formatted,
            override_profile,
            stage="final",
            compare_source=invalid_source,
        )
        override_codes = {item["code"] for item in override_report["issues"]}

        (delivery / "index.md").write_text(
            "internal file list", encoding="utf-8"
        )
        contaminated_delivery_report = audit_delivery(delivery, manifest)

        checks = {
            "formatting": enforce_report["field_codes_preserved"],
            "default_margin_enforced": (
                enforce_report["changes"]["page_margin_sections_formatted"] >= 1
                and margins_match(
                    formatted,
                    {
                        "top": 2.54,
                        "bottom": 2.54,
                        "left": 3.18,
                        "right": 3.18,
                    },
                )
            ),
            "margin_audit_blocks": {
                "PAGE_MARGIN_LEFT",
                "PAGE_MARGIN_RIGHT",
            }.issubset(unformatted_margin_codes),
            "manuscript_audit": manuscript_report["status"] == "pass",
            "default_imrad_audit": (
                manuscript_report["structure"]["enabled"]
                and manuscript_report["structure"]["observed_order"]
                == [
                    "Introduction",
                    "Methods",
                    "Results",
                    "Discussion",
                    "Conclusion",
                ]
            ),
            "structure_boundary_blocks": required_structure_codes.issubset(
                invalid_codes
            ),
            "explicit_format_override": (
                not override_report["structure"]["enabled"]
                and required_structure_codes.isdisjoint(override_codes)
                and override_report["status"] == "pass"
                and margins_match(override_formatted, override_margins)
            ),
            "figure_audit": figure_report["status"] == "pass",
            "delivery_audit": delivery_report["status"] == "pass",
            "citation_limit_blocks": (
                excessive_citation_report["status"] == "fail"
                and any(
                    item["code"] == "CITATION_LIMIT"
                    for item in excessive_citation_report["issues"]
                )
            ),
            "low_dpi_blocks": (
                low_dpi_report["status"] == "fail"
                and any(
                    item["code"] == "FIGURE_DPI"
                    for item in low_dpi_report["issues"]
                )
            ),
            "unregistered_file_blocks": (
                contaminated_delivery_report["status"] == "fail"
                and any(
                    item["code"] == "UNREGISTERED_FILE"
                    for item in contaminated_delivery_report["issues"]
                )
            ),
        }
        return {
            "status": "pass" if all(checks.values()) else "fail",
            "checks": checks,
            "reports": {
                "manuscript": manuscript_report,
                "unformatted_margins": unformatted_margin_report,
                "figures": figure_report,
                "delivery": delivery_report,
                "excessive_citation": excessive_citation_report,
                "invalid_structure": invalid_structure_report,
                "explicit_format_override": override_report,
                "low_dpi": low_dpi_report,
                "contaminated_delivery": contaminated_delivery_report,
            },
        }
    finally:
        shutil.rmtree(temp_root, ignore_errors=True)


def main() -> int:
    result = run_self_test()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
